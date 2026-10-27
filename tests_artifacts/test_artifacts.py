"""Optional report/deck readiness and provenance checks; no provider calls."""

import contextlib
import csv
import io
from pathlib import Path
import os
import socket
import tempfile
import unittest
from unittest.mock import patch

from calendar_pipeline.core import ROOT,read_json,save_json
import tools.build_research_artifacts as artifacts
import tools.prepare_discussion_artifacts as preparation
from tools.finalize_study import readiness
from pptx import Presentation


class ArtifactTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory()
        cls.folder=Path(cls.temp.name)
        (cls.folder/"data").symlink_to(ROOT/"data",target_is_directory=True)
        (cls.folder/"experiments").symlink_to(ROOT/"experiments",target_is_directory=True)
        with patch.object(preparation,"ROOT",cls.folder),patch.object(artifacts,"ROOT",cls.folder),patch.object(socket,"create_connection",side_effect=AssertionError("No network during report build")),contextlib.redirect_stdout(io.StringIO()):
            preparation.build()

    @classmethod
    def tearDownClass(cls):cls.temp.cleanup()

    def test_nonreal_or_partial_study_is_rejected_before_writing_research_rows(self):
        for name,backend,status in [("scripted","scripted-reference","completed"),("partial","mistral","partial")]:
            path=self.folder/name;path.mkdir()
            save_json(path/"metadata.json",{"backend":backend})
            save_json(path/"summary.json",{"run_status":status})
            with patch.object(artifacts,"summarize",side_effect=AssertionError("Must not aggregate rejected results")),self.assertRaises(ValueError):
                artifacts.final_evidence(path)

    def test_readiness_is_local_and_does_not_return_credentials(self):
        fixture="test-only-never-real-credential"
        with patch.dict(os.environ,{"MISTRAL_API_KEY":fixture,"MISTRAL_API_KEY2":fixture,"MISTRAL_API_KEY3":fixture}),patch.object(socket,"create_connection",side_effect=AssertionError("Dry-run must be local")):
            state=readiness(ROOT/"experiments/self_healing_plan_v1.json")
        self.assertEqual(state["planned_fresh_sessions"],160)
        self.assertEqual(state["nominal_logical_requests"],1120)
        self.assertNotIn(fixture,str(state))
        self.assertEqual(state["connectivity"],"not tested by dry-run")

    def test_preparation_chart_rows_are_actual_existing_observations(self):
        with (self.folder/"reports/preparation_chart_rows.csv").open(newline="") as f:rows=list(csv.DictReader(f))
        self.assertEqual(sum(int(r["sessions"]) for r in rows),40)
        self.assertEqual(sum(int(r["candidate_sessions"]) for r in rows),4)
        self.assertEqual(len(rows),10)
        self.assertTrue(all(int(r["sessions"])==4 for r in rows))

    def test_preparation_has_no_fabricated_fresh_result_and_uses_embedded_pdf_fonts(self):
        source=read_json(self.folder/"reports/preparation_report_source.json")
        self.assertIsNone(source["fresh_defense_results"])
        self.assertIsNone(source["human_confirmed_results"])
        pdf=(self.folder/"reports/preparation_report.pdf").read_bytes()
        self.assertTrue(pdf.startswith(b"%PDF-"))
        self.assertIn(b"DejaVuSans",pdf)
        text=(self.folder/"reports/preparation_report.md").read_text()
        self.assertIn("does not report unexecuted defense results",text)

    def test_ten_editable_slides_have_notes_and_fit_slide_bounds(self):
        prs=Presentation(self.folder/"slides/preparation.pptx")
        self.assertEqual(len(prs.slides),10)
        for slide in prs.slides:
            self.assertTrue(slide.notes_slide.notes_text_frame.text.strip())
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left,0);self.assertGreaterEqual(shape.top,0)
                self.assertLessEqual(shape.left+shape.width,prs.slide_width+100)
                self.assertLessEqual(shape.top+shape.height,prs.slide_height+100)


if __name__=="__main__":unittest.main()
