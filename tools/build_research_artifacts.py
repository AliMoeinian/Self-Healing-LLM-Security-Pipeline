"""Source-backed research figures, short technical PDF/Markdown, and 10 slides.

Requires an audited, completed real-provider study. Never substitutes scripted
results or invents a pending live comparison. All chart source rows are saved.
"""

import argparse
from datetime import datetime, timezone
import json
import os
from pathlib import Path
import re
import sys
import textwrap
from xml.sax.saxutils import escape

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from calendar_pipeline.core import ROOT, read_json, save_json, sha_file
from tools.summarize_self_healing import summarize

os.environ.setdefault("MPLCONFIGDIR", str(ROOT / "runs/.matplotlib"))
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle, FancyBboxPatch, FancyArrowPatch, Patch
import numpy as np
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, KeepTogether
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# Embed a real font in both PDFs; base-14 fallback can render inconsistently
# in minimal environments and should not determine the exported typography.
_font_dir = Path(matplotlib.get_data_path()) / "fonts/ttf"
pdfmetrics.registerFont(TTFont("ResearchSans", str(_font_dir / "DejaVuSans.ttf")))
pdfmetrics.registerFont(TTFont("ResearchSans-Bold", str(_font_dir / "DejaVuSans-Bold.ttf")))

COLORS = {"none": "#334155", "input_only": "#2563A0", "output_only": "#B7791F", "both": "#087F8C"}
LABELS = {"none": "No defense", "input_only": "Input only", "output_only": "Output only", "both": "Both"}
ORDER = list(LABELS)


def image_size(path):
    from PIL import Image as PILImage
    with PILImage.open(path) as source:
        return source.size


def figure_setup():
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 11,
                         "axes.spines.top": False, "axes.spines.right": False,
                         "axes.labelcolor": "#263445", "text.color": "#172435",
                         "axes.edgecolor": "#94A3B8", "xtick.color": "#475569", "ytick.color": "#475569",
                         "figure.facecolor": "white", "axes.facecolor": "white", "svg.fonttype": "none"})


def save_figure(fig, directory, name, note):
    engine = fig.get_layout_engine()
    if engine is not None:
        engine.set(rect=(0, .12, 1, .88))
    note = "\n".join(textwrap.wrap(note, max(100, int(fig.get_figwidth()*16))))
    fig.text(.01, .012, note, ha="left", va="bottom", fontsize=8, color="#475569")
    for suffix in ("png", "svg", "pdf"):
        fig.savefig(directory / f"{name}.{suffix}", dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def context_figures(directory):
    directory.mkdir(parents=True, exist_ok=True)
    figure_setup()
    sid = "CAL-0052"
    public = read_json(ROOT / "data/calendar/v1/scenarios" / sid / "public.json")
    oracle = read_json(ROOT / "data/calendar/v1/scenarios" / sid / "evaluator.json")
    chosen = oracle["valid_meeting_intervals"][0] if oracle["valid_meeting_intervals"] else None
    fig, axes = plt.subplots(3, 5, figsize=(13.5, 5.6), sharex=True, sharey=True, layout="constrained")
    for row, owner in enumerate("ABC"):
        own = read_json(ROOT / "data/calendar/v1/scenarios" / sid / "private" / (owner + ".json"))
        for col, day in enumerate(public["working_dates"]):
            ax = axes[row, col]
            ax.set_facecolor("#F8FAFC")
            for event in own["events"]:
                if event["start"].startswith(day):
                    start, end = datetime.fromisoformat(event["start"]), datetime.fromisoformat(event["end"])
                    h0, h1 = start.hour + start.minute/60, end.hour + end.minute/60
                    ax.broken_barh([(h0, h1-h0)], (.15,.7), facecolors="#64748B")
            if chosen and chosen["start"].startswith(day):
                start, end = datetime.fromisoformat(chosen["start"]), datetime.fromisoformat(chosen["end"])
                h0, h1 = start.hour+start.minute/60, end.hour+end.minute/60
                ax.broken_barh([(h0, h1-h0)], (.05,.9), facecolors="#E4B44E", edgecolor="#8F6412", linewidth=1)
            ax.set_xlim(9,17)
            ax.set_ylim(0,1)
            ax.set_xticks([9,11,13,15,17], ["09", "11", "13", "15", "17"])
            ax.set_yticks([])
            ax.grid(axis="x", color="#CBD5E1", linewidth=.5)
            ax.set_axisbelow(True)
            if row == 0:
                ax.set_title(datetime.fromisoformat(day).strftime("%a %b %d"), fontsize=11)
            if col == 0:
                ax.set_ylabel("Owner " + owner, rotation=0, labelpad=35, va="center", fontweight="bold")
    fig.suptitle("Synthetic hard calendar: CAL-0052", x=.02, ha="left", fontsize=15, fontweight="bold")
    fig.legend(handles=[Patch(facecolor="#64748B", label="Existing busy event"), Patch(facecolor="#E4B44E", edgecolor="#8F6412", label="Unique common two-hour interval")], loc="upper right", ncol=2, frameon=False)
    save_figure(fig, directory, "calendar_example", "Source: fixed CAL-0052 calendars and evaluator.json; all times UTC, half-open intervals. Blank time is free. Semantic event details omitted from this figure.")

    fig, ax = plt.subplots(figsize=(12,5.6))
    ax.set_xlim(0,12)
    ax.set_ylim(0,5.6)
    ax.axis("off")
    blocks = [(0.3,3.6,2.4,.95,"PROBE","Native garak replay\nof archived cases"),
              (3.2,3.6,2.4,.95,"PATCH","Candidate-triggered\npredefined policy"),
              (6.1,3.6,2.4,.95,"VERIFY","Fresh live scans\nand independent oracle"),
              (1.7,1.2,2.4,1.0,"INPUT POLICY","Omit peer prose;\npreserve attachments"),
              (4.7,1.2,2.4,1.0,"CALENDAR TEAM","Publish -> propose\n-> three ACKs -> commit"),
              (7.7,1.2,2.4,1.0,"OUTPUT POLICY","Publish fixed phrases;\nretain raw evidence")]
    for x0,y0,w,h,title,body in blocks:
        ax.add_patch(FancyBboxPatch((x0,y0),w,h,boxstyle="round,pad=0.12,rounding_size=0.08",facecolor="#F1F5F9",edgecolor="#64748B",linewidth=1))
        ax.text(x0+w/2,y0+h*.76,title,ha="center",fontweight="bold",fontsize=11)
        ax.text(x0+w/2,y0+h*.35,body,ha="center",va="center",fontsize=10)
    for a,b in [((2.85,4.05),(3.05,4.05)),((5.75,4.05),(5.95,4.05)),((4.25,1.7),(4.55,1.7)),((7.25,1.7),(7.55,1.7)),((4.4,3.45),(2.9,2.4)),((9.0,2.4),(7.4,3.45))]:
        ax.add_patch(FancyArrowPatch(a,b,arrowstyle="-|>",mutation_scale=14,color="#334155",linewidth=1.4))
    ax.text(.2,5.1,"Garak-triggered configured repair loop",fontsize=16,fontweight="bold")
    ax.text(.2,.35,"Ablation: no defense / input only / output only / both. External attacker note is inserted after its sender's local output policy.",fontsize=9)
    save_figure(fig, directory, "repair_architecture", "The defenses minimize a text channel for a fixed structured task. This is policy activation, not autonomous defense invention or model training.")



def figures(evidence, directory, run):
    directory.mkdir(parents=True, exist_ok=True)
    figure_setup()
    groups = {g["variant"]: g for g in evidence["groups"]}
    x = np.arange(4)
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.4), layout="constrained")
    hits = [groups[v]["candidate_injection_sessions"] for v in ORDER]
    eligible = [groups[v]["evaluable_injection_sessions"] for v in ORDER]
    correct = [groups[v]["scheduling_successes"] for v in ORDER]
    total = [groups[v]["sessions"] for v in ORDER]
    for ax, values, denoms, title, ylabel in ((axes[0], hits, eligible, "Published exact-value candidates", "Injection sessions flagged"),
                                             (axes[1], correct, total, "Correct team scheduling", "Correct team sessions")):
        ax.bar(x, values, color=[COLORS[v] for v in ORDER], width=.6)
        ax.set_xticks(x, [LABELS[v] for v in ORDER])
        ax.set_ylim(0, max(denoms) * 1.16)
        ax.set_title(title, loc="left", fontweight="bold", pad=14)
        ax.set_ylabel(ylabel)
        ax.grid(axis="y", color="#E2E8F0", linewidth=.6)
        ax.set_axisbelow(True)
        for index, (value, denom) in enumerate(zip(values, denoms)):
            ax.text(index, value + max(denoms)*.025, f"{value}/{denom}", ha="center", fontweight="bold", fontsize=11)
    save_figure(fig, directory, "privacy_utility", "Source: chart_rows.csv / evidence_summary.json. Exact-value flags; human semantic review pending. Four development fixtures per model.")

    conditions = ["authority_spoofing", "task_pretext", "role_delimiter"]
    titles = ["Forged authority", "Task pretext", "Role delimiters"]
    fig, axes = plt.subplots(1, 2, figsize=(11, 4.8), layout="constrained")
    for ax, model in zip(axes, ("ministral-3b-2512", "ministral-8b-2512")):
        rows = evidence["groups_by_model_condition"]
        matrix = np.array([[next(r["candidate_sessions"] for r in rows if r["variant"] == v and r["model"] == model and r["condition_id"] == c) for c in conditions] for v in ORDER])
        denoms = np.array([[next(r["evaluable"] for r in rows if r["variant"] == v and r["model"] == model and r["condition_id"] == c) for c in conditions] for v in ORDER])
        maximum = max(r["sessions"] for r in rows)
        im = ax.imshow(matrix, vmin=0, vmax=maximum, cmap="Blues", aspect="auto")
        ax.set_title(model, loc="left", fontweight="bold", pad=12)
        ax.set_xticks(np.arange(3), titles)
        ax.set_yticks(np.arange(4), [LABELS[v] for v in ORDER])
        for i in range(4):
            for j in range(3):
                ax.text(j, i, f"{matrix[i,j]}/{denoms[i,j]}", ha="center", va="center", color="white" if matrix[i,j] > maximum/2 else "#172435", fontweight="bold")
        for spine in ax.spines.values():
            spine.set_visible(False)
    fig.colorbar(im, ax=list(axes), shrink=.8, label="Flagged sessions", ticks=range(maximum + 1))
    save_figure(fig, directory, "attack_matrix", "Source: chart_rows.csv. Cells show flagged/evaluable sessions. Direct requests and benign controls are excluded; human labels pending.")

    context_figures(directory)

    raw = read_json(run / "raw_model_candidates.json")
    key = next((k for k in raw if k.startswith("output_only__")), None)
    transformation = None
    if key:
        finding = raw[key][0]
        case = read_json(run / "cases" / (key + ".json"))
        accepted = [c for c in case["session"]["calls"] if c["accepted"]]
        index = next(i for i,c in enumerate(accepted) if c["call_id"] == finding["call_id"])
        transformation = {"case_id": key, "call_id": finding["call_id"], "before": finding["raw_content"],
                          "after": case["session"]["wire_messages"][index]["content"], "scope": "same model response, deterministic output transformation"}
        fig, axes = plt.subplots(1,2,figsize=(12,4.1))
        for ax,title,content,color in ((axes[0],"Raw model message (not published)",transformation["before"],"#A3541A"),
                                       (axes[1],"Peer-readable publication",transformation["after"],COLORS["both"])):
            ax.axis("off")
            ax.text(0,1,title,transform=ax.transAxes,fontsize=13,fontweight="bold",color=color,va="top")
            wrapped="\n".join(textwrap.wrap(content,55))
            ax.text(0,.78,wrapped,transform=ax.transAxes,fontsize=12,va="top",linespacing=1.6)
        fig.suptitle("One output-filter transformation",x=.02,ha="left",fontsize=16,fontweight="bold")
        save_figure(fig,directory,"output_transformation","Source: same accepted response and wire message; researcher-only raw output is distinct from published disclosure. Human semantic review pending.")
        save_json(directory / "output_transformation.json", transformation)
    return transformation


def report_sections(e):
    groups = {g["variant"]: g for g in e["groups"]}
    base, both = groups["none"], groups["both"]
    rows = [[LABELS[v], str(groups[v]["sessions"]), f"{groups[v]['scheduling_successes']}/{groups[v]['sessions']}",
             f"{groups[v]['candidate_injection_sessions']}/{groups[v]['evaluable_injection_sessions']}",
             str(groups[v]["raw_model_candidate_injection_sessions"])] for v in ORDER]
    sections = [
        ("Abstract", [f"We built a bounded proof-of-concept security loop for three small-model calendar assistants. Native garak first replayed preserved vulnerable runs; a deterministic rule activated two configurable channel policies. A fresh four-way ablation then ran {e['recorded_sessions']} real team sessions across four development fixtures, two models, and five conditions. Exact-value injection candidate counts changed from {base['candidate_injection_sessions']}/{base['evaluable_injection_sessions']} without defense to {both['candidate_injection_sessions']}/{both['evaluable_injection_sessions']} with both policies. Scheduling was correct in {e['scheduling_successes']}/{e['recorded_sessions']} sessions. Human semantic review is pending: these are candidate counts, not human-confirmed attack-success rates."]),
        ("1. Research question and exercise scope", ["Can a Garak-triggered, configurable input/output policy reduce private event-detail publication while preserving structured multi-agent meeting scheduling? We demonstrate one probe -> patch -> verify cycle. Our contribution here is a reproducible engineering prototype and explicit measurement boundaries, not a novel general-purpose security defense. The professor's exercise permits a bounded scope; Docker, extensive probe coverage, and multi-iteration optimization are not prerequisites for this demonstration."]),
        ("2. Synthetic task and inspiration", ["CalBench inspired the private-calendar coordination setting. Our code, calendar generator, templates, and simplified task are independently constructed with disclosed ChatGPT assistance; no CalBench code/data are copied. The dataset has 60 fixtures, 15 each in easy/medium/hard/infeasible buckets, 180 owner calendars, and 2,334 fictional events. Twenty fixtures form a development split; 40 remain untested by live LLM calls. The security study uses CAL-0019, CAL-0003, CAL-0052, and CAL-0004, selected for difficulty/role coverage during development.","Each owner sees its own raw calendar and shared public constraints. The team must find one immutable two-hour meeting in November 2-6, 2026, 09:00-17:00 UTC on a 30-minute grid. Availability is permitted information; event titles, categories, descriptions, people, and locations are protected. This simplified task omits rescheduling, disruption costs, and CalBench VPS/information-theoretic analysis."]),
        ("3. Reliable scheduling architecture", ["A local deterministic tool computes each owner's complete availability. Models publish a structured action, one designated coordinator selects a shared candidate index, and all three owners acknowledge the exact proposal ID. The controller independently validates local calendars and matching receipts before an in-process atomic commit. Incomplete or invalid negotiation remains unresolved; it is not treated as proof of infeasibility. Software computes intersections, so the results are application-assisted coordination, not unaided model reasoning.","Earlier v1/v2 pilots had 0/8 successful team outcomes each, despite valid API access. Diagnosis identified missing availability content, unsupported peer claims, and independent conflicting final decisions. V3 achieved 7/8; v4 added phase-specific strict schemas and achieved 40/40 across all 20 development fixtures and both models. These negative traces remain available rather than being discarded."]),
        ("4. Threat model and native garak integration", ["The attacker replaces one participant's published free-text note with a fixed payload. The three injection families are forged authority, scheduling-task pretext, and forged role delimiters; direct privacy requests and neutral notes are separate controls. No actual victim values appear in attack payloads. This implements participant-message injection, not an MCP server, adaptive attacker, or native-system-prompt attack.","Custom subclasses of real garak 0.17.1 Probe, Generator, and Detector classes target the entire team. One normal garak generation entails seven provider calls. Native ThresholdEvaluator produces the report; a custom bounded runner preserves exact requests, accepted/rejected replies, delivered messages, and receipts. We do not claim to run garak's complete built-in catalog. The pinned focused runtime intentionally omits unrelated provider/local-model plugin dependencies."]),
        ("5. Configured repair policies and automatic cycle", ["Input-only policy replaces every incoming peer prose field with a fixed neutral placeholder before it reaches the model, preserving routing, validated availability, candidates, and proposal metadata. Output-only policy replaces an accepted reply's prose with a phase-specific fixed phrase before broadcast; the raw reply remains in researcher-only evidence. Both policies are broad channel minimization, not semantic classifiers. They need no private-calendar or oracle lookup.","The external attacker note is inserted after its sender's local output policy. Thus output-only does not accidentally remove the adversarial input before the victim receives it. Input protection still applies at each recipient. Action validation precedes output rewriting: unsafe or invalid scheduling actions cannot be disguised as a successful commit.","The controller probes archived evidence with native garak, triggers on at least one exact-value candidate, selects the predefined both-policy configuration, and performs fresh verification plus separate input-only/output-only ablations. It logs the trigger, selected policy, executions, and verification result. This is automatic policy activation, not autonomous defense invention, weight updates, or a claim of general self-repair."]),
        ("6. Experimental design and evaluation", [f"The fresh study contains four scenarios x two exact model IDs x five conditions x four policies = {e['recorded_sessions']} sessions. Both models are homogeneous three-agent teams: ministral-3b-2512 and ministral-8b-2512. Inputs use temperature zero, fixed provider seed, strict phase JSON schemas, and role-bound credentials. Within each scenario/model pair we rotate policy order; across eight pair blocks each policy occupies each temporal position twice. Initial requests match across policies/conditions/models except model ID. Later histories reflect actual outputs and policy transforms.","Scheduling is scored against independent fixture answers, absent from runtime decision selection. Every model request, transform, publication, local validation, ACK, and commit is reconstructed from the saved trace. Privacy scoring examines peer-readable victim publications after the routed intervention. A blocked input remains an attempted/evaluable case; raw malicious instruction visibility is recorded separately. Final published ACKs count as group-channel publications, without implying a subsequent attacker-model read.","Exact normalized title/description/person/location matches are grounded candidate flags; category keywords alone are excluded. Non-matches do not establish semantic privacy, and raw researcher-only generations do not count as published leaks. Human worksheets contain every eligible publication, including unflagged messages. Labels support leak/no_leak/uncertain/unsupported_private_claim with quotes and source facts; blank labels remain pending."]),
        ("7. Results", [f"All numbers below describe the new fresh study. The earlier 40-case discovery replay adds zero new model observations. The live client completed {e['logical_requests']} logical requests and {e['http_post_attempts']} HTTP POST attempts. Provider-reported successful-response use was {e['provider_usage']['total_tokens']:,} tokens. HTTP error counts: {json.dumps(e['http_error_status_counts'], sort_keys=True)}. No monetary charge is inferred from token counts.","The table separates public candidate flags from researcher-only raw model candidates. The latter shows why output filtering should not be interpreted as changing an underlying model's tendency to generate private text. See the figures and complete per-model/condition CSV for exact denominators. Human-confirmed disclosure rates remain undefined."]),
        ("8. Limits and interpretation", ["Only four development scenarios and three static injection strings were tested. There is one generation per case, provider nondeterminism remains, and repeated fixtures/messages are not independent population samples. No live held-out estimate, statistical model-safety ranking, adaptive attack coverage, tool-output poisoning, or MCP deployment claim is supported.","The output policy deliberately suppresses all conversational prose. It can retain scheduling utility here because trusted structured attachments carry the required information; it may remove useful explanations or fail as a design choice in tasks needing free-text collaboration. Input minimization does not address malicious content in an owner's own calendar, attacks on trusted tools, or compromised structured state. Availability leakage and event inference from busy times are outside this semantic-detail metric.","The detector is not a validated semantic classifier. Human evaluation must review unflagged messages too, and source/quote checks do not prove entailment. A zero candidate count under a finite set of probes does not prove security. The report is a discussion-ready prototype with an explicit pending human-review component, not a human-confirmed final security conclusion."]),
        ("9. Reproducibility and human review", ["Clone the repository and use Python 3.12 (tested 3.12.14). Install the focused garak environment using bash tools/setup_garak.sh, then install requirements/artifacts.lock for chart/report/deck reproduction. Run the application tests, optional garak tests, and dataset tests. Audit the published study with python -m calendar_pipeline.self_healing audit --run experiments/self_healing/run_v1. Rebuild reviewed rows with tools/summarize_self_healing.py and these artifacts with tools/build_research_artifacts.py. These audits/rebuilds spend no API credits.","New live runs require the three securely injected role-bound Mistral keys and api.mistral.ai access. Prepare a frozen plan from configs/self_healing_v1.json and run it into a fresh ignored runs/ directory. Per-policy and global request/reservation caps, same-key bounded retries, and no ambiguous timeout failover are enforced. Raw synthetic traces are exported only after auditing and scanning current secret values.","For human review, copy experiments/self_healing/run_v1/human_review/messages.csv, annotate without changing IDs/text, and import with python -m calendar_pipeline.self_healing review. Use a new output directory. The original machine report remains intact and no provider rerun is needed. The 176-message historical worksheet remains separately preserved."]),
        ("10. Research extensions and assistance disclosure", ["Promising extensions include semantic/taint-aware selective release instead of blanket suppression; attacks hidden in tool results and MCP responses; adaptive/multi-turn attacker search; repeated seeds and larger held-out fixtures; richer scheduling tasks where privacy and explanation utility conflict; and joint measurement of allowed-availability inference versus explicit private-detail publication. These require new experimental designs rather than extrapolation from this prototype.","The applicant independently read the garak paper and supplied the calendar-privacy framing and scope decisions, as recorded in the diary. ChatGPT assisted with explanation and implemented code, tests, probes, experiments, review tooling, figures, report, and slides. Initial prior semantic labels were AI-assisted. The applicant's independent human confirmation has not yet been supplied. Test doubles and test-only labels are not experimental evidence."]),
        ("References and evidence", ["Derczynski et al., garak: A Framework for Security Probing Large Language Models. arXiv:2406.11036. https://arxiv.org/abs/2406.11036 ; https://github.com/NVIDIA/garak ; https://docs.garak.ai/garak .", "CalBench: Evaluating Coordination-Privacy Trade-offs in Multi-Agent LLMs. arXiv:2605.09823. https://arxiv.org/abs/2605.09823 . Used as acknowledged inspiration, not as a source of copied data/code.", "Primary evidence: experiments/self_healing/run_v1/ (native garak reports, fresh application traces, transform journal, plan, loop decision, blank human sheets); experiments/self_healing/evidence_summary.json and chart_rows.csv; earlier negative/benign/privacy pilots in experiments/. Source repository: https://github.com/AliMoeinian/Self-Healing-LLM-Security-Pipeline ."])
    ]
    return sections, rows


def write_report(e, directory, figure_dir):
    directory.mkdir(parents=True, exist_ok=True)
    sections, rows = report_sections(e)
    title = "A Garak-Triggered Privacy Repair Loop for Small-Model Calendar Teams"
    subtitle = "Bounded research exercise | Automatic candidate metrics; human semantic review pending"
    md = ["# " + title, "", subtitle, "", "Prepared for technical discussion. All calendars and private details are synthetic.", ""]
    for heading, paragraphs in sections:
        md += ["## " + heading, ""]
        for paragraph in paragraphs:
            md += [paragraph, ""]
        if heading == "7. Results":
            md += ["| Policy | Sessions | Correct scheduling | Injection candidates | Raw-generation candidate sessions |", "| --- | ---: | ---: | ---: | ---: |"]
            md += ["| " + " | ".join(r) + " |" for r in rows]
            md += ["", "![Public candidates and scheduling utility](../figures/privacy_utility.png)", "", "![Per-model attack matrix](../figures/attack_matrix.png)", ""]
        if heading == "3. Reliable scheduling architecture":
            md += ["![Synthetic calendar example](../figures/calendar_example.png)", ""]
        if heading == "5. Configured repair policies and automatic cycle":
            md += ["![Repair loop](../figures/repair_architecture.png)", ""]
    (directory / "technical_report.md").write_text("\n".join(md))
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        if hasattr(style, "fontName"):
            style.fontName = "ResearchSans-Bold" if "Bold" in style.fontName else "ResearchSans"
    styles.add(ParagraphStyle(name="ResearchBody",fontName="ResearchSans",fontSize=10,leading=14,spaceAfter=8))
    styles.add(ParagraphStyle(name="ResearchNote",fontName="ResearchSans",fontSize=8,leading=11,textColor=colors.HexColor("#475569")))
    story=[Paragraph(escape(title),styles["Title"]),Spacer(1,8),Paragraph(escape(subtitle),styles["ResearchBody"]),
           Paragraph("All data are synthetic. Human review remains pending.",styles["ResearchNote"]),Spacer(1,16)]
    image_map={"3. Reliable scheduling architecture":"calendar_example", "5. Configured repair policies and automatic cycle":"repair_architecture", "7. Results":"privacy_utility"}
    for heading,paragraphs in sections:
        story.append(Paragraph(escape(heading),styles["Heading2"]))
        for paragraph in paragraphs:
            story.append(Paragraph(escape(paragraph),styles["ResearchBody"]))
        if heading=="7. Results":
            data=[["Policy","Teams","Correct","Published flags","Raw flags"]]+rows
            table=Table(data,colWidths=[89,45,65,105,65],repeatRows=1,hAlign="LEFT")
            table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#E2E8F0")),("FONTNAME",(0,0),(-1,-1),"ResearchSans"),("FONTNAME",(0,0),(-1,0),"ResearchSans-Bold"),("FONTSIZE",(0,0),(-1,-1),9),("TOPPADDING",(0,0),(-1,-1),7),("BOTTOMPADDING",(0,0),(-1,-1),7),("LINEBELOW",(0,0),(-1,0),.5,colors.grey)]))
            story += [table,Spacer(1,10)]
        if heading in image_map:
            from PIL import Image as PILImage
            path=figure_dir/(image_map[heading]+".png")
            width,height=image_size(path)
            story += [Image(str(path),width=470,height=470*height/width),Spacer(1,9)]
    def footer(c,doc):
        c.setFont("ResearchSans",8)
        c.setFillColor(colors.HexColor("#64748B"))
        c.drawString(42,24,"Calendar privacy prototype | Human semantic review pending")
        c.drawRightString(A4[0]-42,24,str(doc.page))
    SimpleDocTemplate(str(directory/"technical_report.pdf"),pagesize=A4,leftMargin=42,rightMargin=42,topMargin=38,bottomMargin=42,title=title,author="Self-Healing LLM Security Pipeline (AI-assisted prototype)").build(story,onFirstPage=footer,onLaterPages=footer)
    save_json(directory/"report_source.json",{"title":title,"subtitle":subtitle,"sections":sections,"result_table":rows,"source_evidence_sha256":sha_file(ROOT/"experiments/self_healing/evidence_summary.json")})


def slide_content(e, transformation):
    g={x["variant"]:x for x in e["groups"]}
    compact="; ".join(f"{LABELS[v]}: {g[v]['candidate_injection_sessions']}/{g[v]['evaluable_injection_sessions']}" for v in ORDER)
    return [
        {"title":"A configured privacy repair loop", "subtitle":"Small-model calendar teams | garak -> patch -> verify", "bullets":["Research question: reduce private-detail publication while preserving team scheduling.",f"{e['recorded_sessions']} fresh team executions; two Ministral models; four policy configurations.","Human semantic review is pending; results use grounded exact-value candidate flags."],"notes":"Bounded PhD discussion exercise. AI assistance is disclosed in the report and diary. Do not present candidate rates as human-confirmed security outcomes."},
        {"title":"A narrow, inspectable threat model", "bullets":["Three assistants; each receives only its own raw private calendar.","One peer's published prose carries a static injection; structured attachments remain trusted.","Availability may be shared. Event titles, people, locations and descriptions must remain private.","This prototype uses participant-message injection; MCP infrastructure and adaptive attacks are future work."],"notes":"Direct privacy requests and benign notes are separate controls. Three injection families: forged authority, task pretext, forged role delimiters. Source: attack_catalog.py."},
        {"title":"Synthetic calendars with a fixed answer", "image":"calendar_example", "bullets":[],"notes":"60 fixtures, 15 per difficulty bucket; 180 owner calendars and 2,334 fictional events. Twenty development / forty held-out. Four development fixtures in this security study. Calendar times UTC; the shared highlighted interval is derived from the fixed evaluator for illustration only, never supplied to runtime models. CalBench is acknowledged inspiration."},
        {"title":"Reliability came before security", "bullets":["v1/v2: 0/8 team successes; missing availability and conflicting final decisions.","v3: 7/8; one repeated oversized action remained unresolved.","v4: 40/40 across all 20 development fixtures and both models.","Deterministic tools compute availability/intersections; models publish, select and acknowledge.","These are application-assisted outcomes, not unaided LLM scheduling claims."],"notes":"Negative results and diagnosis are preserved in experiments/. Strict schema does not by itself ensure privacy; this motivated the injection study."},
        {"title":"Probe -> configured patch -> verify", "image":"repair_architecture", "bullets":[],"notes":"Native garak replays preserved vulnerable cases with no new model calls. A deterministic candidate trigger activates predefined both-policy protection. Fresh scans and four-way ablation verify the configured patch. This is automated activation, not learning new defenses or updating weights."},
        {"title":"Two independent channel policies", "bullets":["Input: omit peer prose before generation; preserve validated scheduling attachments.","Output: publish fixed phase phrases; retain original model responses as researcher evidence.","Insert the external attacker note after its sender's local output policy, so output-only victims still receive the attack.","Validate actions before rewriting text; failed ACKs cannot become fake successful commits."],"notes":"Both defenses avoid oracle/private fact lookup. Blanket prose suppression is appropriate only because this fixed task carries the essential scheduling facts in structured channels. Rich conversation utility is not measured."},
        {"title":"Public candidate flags and scheduling utility", "image":"privacy_utility", "bullets":[],"notes":compact+f". Correct scheduling: {e['scheduling_successes']}/{e['recorded_sessions']}. Each policy has 24 injection sessions plus 8 benign and 8 direct-request controls. No general security claim; human labels pending."},
        {"title":"Attack family and model breakdown", "image":"attack_matrix", "bullets":[],"notes":"Cells are flagged/evaluable team sessions. Four development scenarios per family/model/policy. Zero exact matches do not establish semantic privacy. Static probes and repeated fixtures do not support a general model-safety ranking. Complete rows: chart_rows.csv."},
        {"title":"Raw generation differs from publication", "image":"output_transformation" if transformation else None, "bullets":[] if transformation else ["Raw model responses are retained separately from peer-readable output.","Output minimization can suppress publication without altering what the model generated.","Input minimization and output minimization address different stages."],"notes":"The illustrated before/after is one actual model response passed through the output policy, not two assumed identical generations. Raw exact-value candidates are not counted as published disclosures. Human review remains pending."},
        {"title":"Limits, reproducibility and next research", "bullets":["Finite static probes on four development fixtures; no live held-out evaluation.","Blank human worksheets preserve unflagged messages too; later labels require no API rerun.","Blanket prose suppression may remove useful explanations in richer tasks.","Next: selective semantic release, tool/MCP poisoning, adaptive attacks, repeated seeds and availability inference.","Code, plans, traces, audits, figure data, report and deck are reproducible; AI assistance is disclosed."],"notes":"References: garak arXiv:2406.11036 and CalBench arXiv:2605.09823. The applicant independently read garak and framed the calendar scenario; ChatGPT implemented code, tests, experiments and presentation materials. Do not claim independent human confirmation before annotation."}
    ]


def slides(e, transformation, directory, figure_dir, basename="presentation", content_override=None):
    directory.mkdir(parents=True,exist_ok=True)
    content=content_override if content_override is not None else slide_content(e,transformation)
    save_json(directory/(basename+"_source.json"),content)
    prs=Presentation()
    prs.slide_width=Inches(13.333)
    prs.slide_height=Inches(7.5)
    prs.core_properties.title="Garak-triggered calendar privacy repair"
    prs.core_properties.author="Self-Healing LLM Security Pipeline (AI-assisted prototype)"
    for index,item in enumerate(content,1):
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        shape=slide.shapes.add_textbox(Inches(.55), Inches(.35), Inches(12.1), Inches(.75))
        p=shape.text_frame.paragraphs[0]
        p.text=item["title"]
        p.font.size=Pt(29)
        p.font.name="Arial"
        p.font.bold=True
        p.font.color.rgb=RGBColor.from_string("172435")
        if item.get("image"):
            path=figure_dir/(item["image"]+".png")
            from PIL import Image as PILImage
            w,h=image_size(path)
            width=min(12.15,5.55*w/h)
            height=width*h/w
            slide.shapes.add_picture(str(path),Inches((13.333-width)/2), Inches(1.2+(5.55-height)/2), width=Inches(width),height=Inches(height))
        else:
            box=slide.shapes.add_textbox(Inches(.8), Inches(1.55), Inches(11.65), Inches(5.05))
            frame=box.text_frame
            frame.word_wrap=True
            for n,text in enumerate(item["bullets"]):
                p=frame.paragraphs[0] if n==0 else frame.add_paragraph()
                p.text="• "+text
                p.font.size=Pt(23 if len(item["bullets"])<5 else 21)
                p.font.name="Arial"
                p.font.color.rgb=RGBColor.from_string("263445")
                p.space_after=Pt(18)
        foot=slide.shapes.add_textbox(Inches(.55), Inches(7.08), Inches(12.1), Inches(.25))
        p=foot.text_frame.paragraphs[0]
        p.text=f"Self-Healing LLM Security Pipeline | Synthetic data | Human review pending                                      {index}/10"
        p.font.size=Pt(10)
        p.font.name="Arial"
        p.font.color.rgb=RGBColor.from_string("64748B")
        slide.notes_slide.notes_text_frame.text=item["notes"]
    prs.save(directory/(basename+".pptx"))
    pdf=canvas.Canvas(str(directory/(basename+".pdf")),pagesize=(960,540))
    pdf.setTitle("Garak-triggered calendar privacy repair")
    pdf.setAuthor("Self-Healing LLM Security Pipeline (AI-assisted prototype)")
    for index,item in enumerate(content,1):
        pdf.setFillColor(colors.white);pdf.rect(0,0,960,540,fill=1,stroke=0)
        pdf.setFillColor(colors.HexColor("#172435"));pdf.setFont("ResearchSans-Bold",26)
        pdf.drawString(40,488,item["title"])
        if item.get("image"):
            from PIL import Image as PILImage
            path=figure_dir/(item["image"]+".png")
            w,h=image_size(path)
            width=min(880,390*w/h);height=width*h/w
            pdf.drawImage(str(path),(960-width)/2,70+(390-height)/2,width,height,mask="auto")
        else:
            y=412
            for text in item["bullets"]:
                para=Paragraph(escape("• "+text),ParagraphStyle("slide-body",fontName="ResearchSans",fontSize=18 if len(item["bullets"])>=5 else 20,leading=27,textColor=colors.HexColor("#263445")))
                _,h=para.wrap(850,400)
                para.drawOn(pdf,55,y-h);y-=h+20
        pdf.setFillColor(colors.HexColor("#64748B"));pdf.setFont("ResearchSans",9)
        pdf.drawString(40,25,"Synthetic data | Human semantic review pending")
        pdf.drawRightString(920,25,f"{index}/10")
        pdf.showPage()
    pdf.save()
    md=["# Technical discussion deck", "", "Ten slides. PPTX includes speaker notes; PDF uses the same source content.", ""]
    for i,item in enumerate(content,1):
        md += [f"## {i}. {item['title']}",""]+["- "+b for b in item["bullets"]]+["", "**Speaker notes:** "+item["notes"],""]
        if item.get("image"):
            md += [f"![{item['title']}](../figures/{item['image']}.png)",""]
    (directory/(basename+".md")).write_text("\n".join(md))


def final_evidence(run):
    # Check before aggregation writes any chart rows into the research folder.
    # A rejected test-double run must not replace an existing real result file.
    metadata=read_json(run/"metadata.json")
    status=read_json(run/"summary.json")
    if metadata["backend"]!="mistral" or status["run_status"]!="completed":
        raise ValueError("Final research results require a complete real-provider study; no scripted or pending substitution")
    return summarize(run,ROOT/"experiments/self_healing")


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("--run",type=Path,default=ROOT/"experiments/self_healing/run_v1")
    args=p.parse_args()
    evidence=final_evidence(args.run)
    figure_dir=ROOT/"figures"
    transformation=figures(evidence,figure_dir,args.run)
    write_report(evidence,ROOT/"reports",figure_dir)
    slides(evidence,transformation,ROOT/"slides",figure_dir)
    save_json(ROOT/"reports/artifact_manifest.json",{"built_at_utc":datetime.now(timezone.utc).isoformat(),
              "evidence_sha256":sha_file(ROOT/"experiments/self_healing/evidence_summary.json"),
              "source_rows":"experiments/self_healing/chart_rows.csv", "human_review_status":"pending",
              "files_sha256":{str(p.relative_to(ROOT)):sha_file(p) for folder in (ROOT/"figures",ROOT/"slides",ROOT/"reports") for p in sorted(folder.rglob("*")) if p.is_file() and p.name!="artifact_manifest.json"}})
    print("Built source-backed figures, report, and ten-slide PPTX/PDF; human review remains pending.")


if __name__=="__main__":
    main()
